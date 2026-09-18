import argparse
import json
import sys
import time
from collections.abc import Mapping
from pathlib import Path

from openpyxl import load_workbook
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait


NAVER_CRAWLING_PYTHON = Path(
    r"C:\Users\user1\.codex\worktrees\ec47\naver_crawling\scripts\python"
)
sys.path.insert(0, str(NAVER_CRAWLING_PYTHON))

from smartplace_browser import create_uc_driver, naver_javascript_login


DEFAULT_WORKBOOK = Path(
    r"C:\Users\user1\Downloads\레포트사전확인요청_전략영업팀_SCM.xlsx"
)
PROFILE_DIR = Path(__file__).resolve().parents[1] / ".profiles" / "naver-smartstore"
COMMERCE_LOGIN = (
    "https://accounts.commerce.naver.com/login"
    "?url=https%3A%2F%2Fsell.smartstore.naver.com%2F%23%2Flogin-callback"
)
DASHBOARD_API = (
    "https://sell.smartstore.naver.com/api/v1/sellers/"
    "dashboards/pay/order-delivery"
)
COUNTER_KEYS = (
    "newOrderCases",
    "deliveryPreparingCases",
    "deliveringCases",
    "deliveredCases",
    "purchaseCompleteCases",
)


def summarize_dashboard(payload):
    if not isinstance(payload, Mapping):
        raise TypeError("dashboard response must be an object")
    return {key: payload.get(key) for key in COUNTER_KEYS}


def read_credentials(workbook_path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["01.계정 관리"]
        for site, _, user_id, password in sheet.iter_rows(
            min_row=3, min_col=3, max_col=6, values_only=True
        ):
            if site == "네이버 스마트스토어":
                if not user_id or not password:
                    break
                return str(user_id), str(password)
    finally:
        workbook.close()
    raise RuntimeError("스마트스토어 로그인 정보를 찾지 못했습니다.")


def click_naver_id_login(driver):
    button = WebDriverWait(driver, 20).until(
        lambda current: next(
            (
                element
                for element in current.find_elements(By.TAG_NAME, "button")
                if "네이버 아이디로" in element.text and element.is_displayed()
            ),
            None,
        )
    )
    parent_handle = driver.current_window_handle
    button.click()

    WebDriverWait(driver, 30).until(
        lambda current: (
            "sell.smartstore.naver.com" in current.current_url
            or len(current.window_handles) > 1
        )
    )
    if len(driver.window_handles) > 1:
        popup_handle = next(
            handle for handle in driver.window_handles if handle != parent_handle
        )
        driver.switch_to.window(popup_handle)
        WebDriverWait(driver, 30).until(
            lambda current: (
                "nid.naver.com" not in current.current_url
                or len(current.window_handles) == 1
            )
        )
        if parent_handle in driver.window_handles:
            driver.switch_to.window(parent_handle)

    WebDriverWait(driver, 30).until(
        lambda current: "sell.smartstore.naver.com" in current.current_url
    )


def call_dashboard_api(driver):
    return driver.execute_async_script(
        """
        const url = arguments[0];
        const done = arguments[arguments.length - 1];
        fetch(url, { method: 'GET', credentials: 'include' })
          .then(async response => ({
            status: response.status,
            ok: response.ok,
            body: await response.json().catch(() => null),
          }))
          .then(done)
          .catch(error => done({ status: 0, ok: false, error: String(error) }));
        """,
        DASHBOARD_API,
    )


def run(workbook_path, port):
    user_id, password = read_credentials(workbook_path)
    driver = create_uc_driver(
        port=port,
        profile=PROFILE_DIR,
        version_main=150,
        headless=False,
    )
    try:
        naver_javascript_login(driver, user_id, password)
        user_id = password = None

        driver.get(COMMERCE_LOGIN)
        click_naver_id_login(driver)
        WebDriverWait(driver, 30).until(
            lambda current: current.execute_script(
                "return document.readyState"
            ) == "complete"
        )
        time.sleep(1)

        result = call_dashboard_api(driver)
        if not result or not result.get("ok"):
            status = result.get("status") if result else "unknown"
            raise RuntimeError(f"내부 API 호출 실패: HTTP {status}")
        print(
            json.dumps(
                {
                    "status": result["status"],
                    "counters": summarize_dashboard(result["body"]),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    finally:
        user_id = password = None
        driver.quit()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--port", type=int, default=9226)
    args = parser.parse_args()
    run(args.workbook, args.port)


if __name__ == "__main__":
    main()
